"""Gantt KJA. Requiere openpyxl; Excel COM se usa únicamente para recalcular y revisar."""
from pathlib import Path
from datetime import date, datetime, timedelta
from calendar import monthrange
import subprocess, json, re
from collections import Counter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.properties import CalcProperties

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'outputs'/'gantt-kja-20260922'
OUT.mkdir(parents=True,exist_ok=True)
CUT=date(2026,9,22)
NAVY='16324F'; BLUE='DCE9F7'; GRAY='F0F2F5'; WHITE='FFFFFF'; INK='172B40'
COLORS=['F4CAB2','C9EDF6','ECD0EA','D9E8CC','DAD7F4','F9E8B1']
THIN=Side(style='thin',color='B6C0CA')
GIT=['git','-c',f'safe.directory={ROOT.as_posix()}']
log=subprocess.check_output(GIT+['log','--date=iso-strict','--pretty=format:%h|%ad|%an|%s'],cwd=ROOT,text=True,encoding='utf-8')
commits={}
for line in log.splitlines():
    sha,dt,author,subject=line.split('|',3)
    commits[sha]={'hash':sha,'date':dt[:10],'author':author,'subject':subject,'timestamp':dt}
tasks=[]
def task(sprint,hu,title,area,owner,start,end,kind,criterion,rf='',deps='',sources='',days=None):
    ident=f'T-{len(tasks)+1:03d}'
    tasks.append(dict(id=ident,sprint=sprint,hu=hu,title=title,area=area,owner=owner,start=start,end=end,
                      kind=kind,criterion=criterion,rf=rf,deps=deps,sources=sources,days=days or []))
    return ident
def historical(sprint,hu,title,hashes,rf,criterion):
    evidence=[commits[h] for h in hashes]
    days=sorted(set(x['date'] for x in evidence))
    return task(sprint,hu,title,'Transversal',', '.join(sorted(set(x['author'] for x in evidence))),days[0],days[-1],
                'Registro Git',criterion,rf,sources='; '.join(f"{x['hash']} | {x['timestamp']} | {x['subject']}" for x in evidence),days=days)

H={
'HU-G01':'Como colaborador, quiero acceder al portal y consultar mi asistencia.',
'HU-G02':'Como colaborador, quiero registrar mi jornada según modalidad y ubicación.',
'HU-G03':'Como colaborador, quiero gestionar mi perfil y mis solicitudes.',
'HU-G04':'Como colaborador, quiero entregar evidencias y conocer el cierre de mi jornada.',
'HU-G05':'Como responsable, quiero asignar trabajo y supervisar el cumplimiento.',
'HU-G06':'Como usuario, quiero comunicarme con mi equipo y consultar avisos.',
'HU-G07':'Como Marketing, quiero revisar y preparar publicaciones autorizadas.',
'HU-G08':'Como Dirección, quiero consultar reportes y métricas explicables.',
'HU-G09':'Como equipo, queremos requisitos y planificación trazables.',
'HU-G10':'Como responsable de área, quiero actualizar equipo, horarios y necesidades.',
'HU-G11':'Como usuario, quiero mejoras de calidad y flujos confiables.',
'HU-G12':'Como Dirección, quiero validar la adopción y priorizar el siguiente periodo.'}

historical(0,'HU-G01','Acceso, permisos e integración inicial del dashboard',['3788de4','d189536'],'RF-001–RF-014, RF-049–RF-070','Consultar el commit y la guía de despliegue; no deducir la duración de desarrollo.')
task(0,'HU-G01','Hito: despliegue inicial del portal documentado','Transversal','Por confirmar','2026-08-28','2026-08-28','Despliegue documentado',
     'La guía declara producción el 28/08/2026; no se vuelve a verificar el entorno remoto.','RF-062',sources='docs/dashboard-base.md, encabezado Estado; commit d189536',days=['2026-08-28'])
historical(0,'HU-G03','Fotografía privada, avatar y mejora del ingreso',['0f25307','24e8bb1','fd0bfe6'],'RF-127, RNF-024','Revisar archivos incorporados y separar mejora de código de medición de rendimiento.')
historical(1,'HU-G02','Marcado seguro, geocerca y mapa de oficina',['952eb39','6b8b602','6989792'],'RF-025–RF-040, RF-126, RF-130–RF-131','Cambios registrados el 01/09; permisos y geocerca requieren su prueba de aceptación.')
historical(1,'HU-G01','Acceso móvil y experiencia del panel administrativo',['f7d42dc','474f477','11047df'],'RF-016–RF-024, RF-064–RF-070','Constancia de cambios en acceso y gestión del 01–02/09.')
historical(1,'HU-G03','Solicitudes personales, fotos y días libres',['02439f8','6e3e60e'],'RF-127–RF-129','Verificar flujo personal y resolución administrativa en un entorno autorizado.')
historical(1,'HU-G04','Cierre diario con evidencias y salida',['4c25c26'],'RF-132–RF-139','Registro integrado el 07/09; esta fecha no representa el inicio de programación.')
historical(1,'HU-G04','Corrección de evidencias y comprobante de Facebook',['9bb0c8c','e61a507','a5aec29','270d345'],'RF-137, RF-144–RF-147','Se registran corrección, comprobante y acceso a compartir por acción del usuario.')
historical(1,'HU-G05','Revisión, observaciones y notificaciones de evidencias',['8b87a46','2752fbf','306e00e'],'RF-152–RF-156','Días observados: 07, 08 y 09/09; no implica dedicación exclusiva durante todo el intervalo.')
historical(1,'HU-G04','Agenda Facebook independiente y vencimientos',['843ea64','15244f2'],'RF-144','Verificar obligación independiente de asistencia y conservación de salida.')
historical(1,'HU-G05','Gestión operativa y roles',['ad6304a'],'RF-114–RF-125, RF-155','Revisar roles y ámbito del área; el commit no acredita despliegue.')
task(1,'HU-G07','Activación técnica parcial de Marketing documentada','Marketing','Por confirmar','2026-09-12','2026-09-12','Despliegue documentado',
     'Guía registra OPTIONS 200 y POST sin sesión 401; no acredita extracción ni publicación real.','RF-170–RF-175',
     sources='docs/18-publicaciones-marketing.md, Activación pendiente (12/09/2026)',days=['2026-09-12'])
historical(1,'HU-G06','Integración de chat privado y ajustes móviles',['cfdac3d','a324b3e'],'RF-165–RF-169','El código queda registrado el 13/09; no se presupone activación completa del chat.')
historical(1,'HU-G07','Integración del editor de publicaciones',['cfdac3d'],'RF-170–RF-175','Constancia del frontend y función en Git; validación externa pendiente.')
historical(1,'HU-G08','Ranking mensual beta y detalle de métricas',['9fa90d0','be8c48f'],'RF-157–RF-159','Revisar fórmula provisional y simulación local de pesos.')
historical(1,'HU-G02','Corrección de espera y verificación de ubicación',['49f9437','7419211'],'RF-039, RF-130','Probar ubicación válida, denegada e imprecisa; registro de cambios del 14/09.')
historical(1,'HU-G05','Asignaciones con documentos, panel y retiro de archivos',['7419211','d736789','60fb768','8d65580'],'RF-138, RF-148–RF-153','Verificar documentos permitidos, canceladas y recuperación del borrado.')
historical(1,'HU-G06','Chat con alertas, imágenes y sonido',['d192f69','8c0bfdd','cbec54a'],'RF-165–RF-169','Registros del 14 y 15/09; probar comunicación con dos cuentas autorizadas.')
historical(1,'HU-G04','Respeto de jornadas justificadas',['1dd2ad5'],'RF-143','J conserva registros y Facebook solo aplica según agenda.')
historical(1,'HU-G04','Preparación de comprobantes para WhatsApp',['8d7ea4a'],'RF-147','Preparar comprobante sin envío automático a contactos.')
historical(1,'HU-G09','Primera documentación de secuencias administrativas',['2c13fe0'],'RNF-079','Conservar el commit documental como antecedente de la actualización posterior.')
historical(1,'HU-G08','Reportes Facebook por áreas con Excel y PDF',['1cb45a3','271bb95'],'RF-160–RF-163','Registro el 17/09 en Lima; no equivale a aceptación con datos reales.')
historical(1,'HU-G08','Estadísticas del panel y ajuste del periodo de salidas',['b254094'],'RF-157–RF-159','Distinguir ranking de cumplimiento y ranking del PDF por total de Sí.')
historical(1,'HU-G04','Exigir evidencia laboral antes de registrar salida',['ee7555f'],'RF-134–RF-136','Verificar rechazo por requisitos pendientes y conservación de entrada.')
historical(1,'HU-G01','Ajustes de tema, tarjetas y experiencia de equipo',['c31d0b9','accf585','dddf078'],'RNF-038, RNF-041','Días con cambios: 19 y 20/09; pendiente evaluación completa de accesibilidad.')
historical(1,'HU-G06','Comunicados emergentes y carrusel del dashboard',['c2de35f','0d868e6','3fcf72e'],'Por vincular: comunicado/carrusel','Registrar posteriormente su RF y H.U. formal; no forzar una asociación inexistente.')
historical(1,'HU-G04','Excepción individual de cierre y RPE presencial',['5dfcf5a','cd33b15'],'RF-140–RF-143','Aplicar fecha efectiva de política; preservar RPE histórico y reglas de otras personas.')
task(1,'HU-G09','Actualización de RF/RNF, 22 secuencias y matriz','Transversal','Por confirmar','2026-09-22','2026-09-22','Documento local',
     'Existen 175 RF, 81 RNF y 22 secuencias; se verificó estructura, no aceptación funcional.','RNF-079',
     sources='docs/01-requerimientos-portal-asistencia.md; docs/20-diagramas-secuencia-administracion.md; docs/27-matriz-requerimientos.md',days=['2026-09-22'])

def propose(sprint,hu,title,area,owner,start,end,criterion,rf='',deps=''):
    return task(sprint,hu,title,area,owner,start,end,'Propuesta',criterion,rf,deps,sources='Propuesta de planificación al 22/09/2026; fechas y responsables por confirmar')

p1=propose(1,'HU-G09','Revisar alcance, responsables y catálogo vigente de áreas','Transversal','Dirección + Sistemas','2026-09-23','2026-09-24',
           'Acta con áreas vigentes, dueño funcional por área y capacidad del equipo.','RF-079–RF-089, RNF-079')
p2=propose(1,'HU-G09','Formalizar H.U. y casos de uso; vincular al catálogo RF/RNF','Transversal','Sistemas + líderes','2026-09-25','2026-09-28',
           'H.U. priorizadas con criterios; cada propuesta HU-G queda confirmada, corregida o descartada.','RNF-079',p1)
p3=propose(1,'HU-G11','Revisar despliegues pendientes y ejecutar pruebas críticas','Transversal','Sistemas + QA','2026-09-29','2026-09-29',
           'Inventario SQL/Edge/frontend y acta de acceso, geocerca, cierre, excepciones y reportes.','RNF-047, RNF-081',p2)
p4=propose(1,'HU-G09','Revisión y retrospectiva del Sprint 1; acordar octubre','Transversal','Dirección + equipo','2026-09-30','2026-09-30',
           'Backlog priorizado, pendientes visibles y capacidad confirmada para octubre.','RNF-079',p3)

areas=[('Salud ocupacional','2026-10-01','2026-10-02','Validar horarios/modalidad, evidencias y responsable de revisión.'),
('Psicología organizacional','2026-10-05','2026-10-06','Validar responsables, equipo y reglas de seguimiento.'),
('Recursos Humanos','2026-10-07','2026-10-09','Depurar altas/bajas, contratos, solicitudes y fichas.'),
('Marketing','2026-10-12','2026-10-13','Revisar agenda Facebook y permisos de borradores/publicación.'),
('Ingeniería','2026-10-14','2026-10-15','Revisar entregables documentales, asignaciones y revisión.'),
('Diseño gráfico','2026-10-16','2026-10-19','Acordar formatos, versiones y revisión de archivos gráficos.'),
('Clínica','2026-10-20','2026-10-21','Validar acceso mínimo y evidencias laborales sin datos de pacientes.'),
('Audiovisuales','2026-10-22','2026-10-23','Revisar video, límites de archivo y recuperación de cargas.'),
('Reclutamiento','2026-10-26','2026-10-27','Validar altas de colaboradores, solicitudes y revisión; sin incorporar expedientes de candidatos.'),
('Contabilidad','2026-10-28','2026-10-29','Verificar ficha, contrato y reportes laborales; sin planillas ni pagos.')]
previous=p4
area_ids=[]
for area,start,end,focus in areas:
    previous=propose(2,'HU-G10','Revisar y actualizar configuración del área: '+area,area,'Sistemas + responsable del área',start,end,
        focus+' Entregar ficha actualizada, una prueba con el área y hasta tres mejoras priorizadas.','RF-079–RF-090, RF-114–RF-125, RF-148–RF-155',previous)
    area_ids.append(previous)
p5=propose(2,'HU-G09','Revisión y retrospectiva del Sprint 2; priorizar mejoras','Transversal','Dirección + responsables','2026-10-30','2026-10-30',
           'Consolidar las diez áreas, cambios aceptados y propuestas elegidas para noviembre.','RNF-079',previous)

nov=[
('HU-G11','Auditar permisos, sesiones y aislamiento por rol','Sistemas + QA','2026-11-02','2026-11-04','Acta de accesos propios/ajenos; corregir hallazgos críticos y revalidar.','RNF-001–RNF-012, RNF-064'),
('HU-G11','Mejorar estados de error, accesibilidad y uso móvil','Sistemas + Diseño gráfico','2026-11-05','2026-11-10','Cerrar hallazgos priorizados de teclado, contraste, foco y pantallas pequeñas.','RNF-035–RNF-042'),
('HU-G11','Medir y optimizar carga, consultas e imágenes','Sistemas + QA','2026-11-11','2026-11-13','Medición p95 reproducible antes/después; optimizar consultas o archivos según evidencia.','RNF-023–RNF-029'),
('HU-G05','Mejorar asignaciones, revisión y recuperación de archivos','Sistemas + líderes','2026-11-16','2026-11-18','Probar versiones, cancelación y reintentos de Storage con dos áreas piloto.','RF-137–RF-139, RF-148–RF-154, RNF-070'),
('HU-G08','Conciliar reportes y ranking con casos de todas las áreas','Sistemas + Dirección','2026-11-19','2026-11-20','Resultados trazables para virtual, presencial, J, excepciones y Facebook sin jornada.','RF-140–RF-163, RNF-069, RNF-077'),
('HU-G06','Validar chat y notificaciones con cuentas de prueba','Sistemas + responsables','2026-11-23','2026-11-24','Probar privacidad, lectura, imágenes, reintento y cierre de sesión.','RF-165–RF-169, RNF-071, RNF-073'),
('HU-G07','Validar Marketing en entorno de ensayo autorizado','Sistemas + Marketing','2026-11-25','2026-11-27','Probar borradores, cuota, revisión y resultado incierto; una publicación real requiere autorización aparte.','RF-170–RF-175, RNF-075–RNF-076'),
('HU-G09','Revisión y retrospectiva del Sprint 3','Dirección + equipo','2026-11-30','2026-11-30','Demostración de mejoras, defectos pendientes y alcance de aceptación de diciembre.','RNF-079')]
previous=p5
for hu,title,owner,start,end,ca,rf in nov:
    previous=propose(3,hu,title,'Transversal',owner,start,end,ca,rf,previous)

dec=[
('HU-G11','Definir retención, respaldo y procedimiento de recuperación','Dirección + Sistemas','2026-12-01','2026-12-03','Política ratificada y ensayo de recuperación con RPO/RTO medidos.','RNF-015–RNF-017, RNF-030–RNF-034'),
('HU-G12','Preparar capacitación y guías por rol','Sistemas + RR. HH.','2026-12-04','2026-12-07','Guías breves de acceso, marcado, cierre, revisión y soporte.','RNF-035, RNF-044'),
('HU-G12','Aceptación con Salud ocupacional, Psicología y Clínica','QA + responsables de área','2026-12-10','2026-12-11','Un acta por área con escenarios propios y pendientes.','RF-025–RF-045, RF-132–RF-143'),
('HU-G12','Aceptación con RR. HH., Reclutamiento y Contabilidad','QA + responsables de área','2026-12-14','2026-12-15','Validar fichas, permisos, solicitudes y reportes según rol.','RF-079–RF-099, RF-128–RF-129, RF-160–RF-164'),
('HU-G12','Aceptación con Marketing, Ingeniería, Diseño y Audiovisuales','QA + responsables de área','2026-12-16','2026-12-17','Validar agenda, archivos, revisión y flujos permitidos por área.','RF-137–RF-139, RF-144–RF-153, RF-170–RF-175'),
('HU-G11','Corregir hallazgos y ejecutar regresión integral','Sistemas + QA','2026-12-18','2026-12-23','Cerrar fallos críticos; evidenciar compatibilidad y resultados por rol.','RNF-047, RNF-062, RNF-081'),
('HU-G09','Actualizar documentación, matriz y manual de operación','Sistemas + responsables','2026-12-24','2026-12-28','RF/RNF, H.U., CU, secuencias y matriz reflejan lo aceptado; registrar pendientes.','RNF-044, RNF-079'),
('HU-G12','Revisar adopción y priorizar propuestas para enero','Dirección + responsables','2026-12-29','2026-12-30','Indicadores de uso/errores y backlog siguiente sin compromisos no acordados.','RNF-051, RNF-079'),
('HU-G09','Revisión y retrospectiva del Sprint 4; cierre anual','Dirección + equipo','2026-12-31','2026-12-31','Acta de aceptación, pendientes, responsables y propuesta de Sprint 5 enero 2027.','RNF-079')]
for hu,title,owner,start,end,ca,rf in dec:
    previous=propose(4,hu,title,'Transversal',owner,start,end,ca,rf,previous)

# Calendario de trabajo propuesto; no es una certificación del calendario laboral.
excluded={'2026-10-08','2026-12-08','2026-12-09','2026-12-25'}
dates=[]; day=date(2026,8,28)
while day<=date(2026,12,31):
    dates.append(day); day+=timedelta(days=1)
cal={d.isoformat():int(d.weekday()<5 and d.isoformat() not in excluded) for d in dates}
for t in tasks:
    a=date.fromisoformat(t['start']); b=date.fromisoformat(t['end'])
    assert a<=b
    assert (a.month-8 if a.month>=9 else 0)==t['sprint']
    assert a.month==b.month
    t['marked_days']=[d.isoformat() for d in dates if a<=d<=b and (cal[d.isoformat()] if t['kind']=='Propuesta' else d.isoformat() in t['days'])]
    t['duration']=len(t['marked_days'])
    assert t['duration']>0
    if t['kind']=='Propuesta': assert a>CUT
    for dependency in t['deps'].split(','):
        if dependency:
            predecessor=next(x for x in tasks if x['id']==dependency)
            assert predecessor['end']<t['start'],(t['id'],dependency)

wb=Workbook(); wb.remove(wb.active)
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
read=wb.create_sheet('LEEME')
source=wb.create_sheet('TAREAS')
calendar=wb.create_sheet('CALENDARIO')
wsheets={s:wb.create_sheet('ANTECEDENTE' if s==0 else f'SPRINT {s}') for s in range(5)}

def basics(ws):
    ws.sheet_view.showGridLines=False
    ws.sheet_view.zoomScale=80
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.print_options.horizontalCentered=True
    ws.page_margins.left=.22; ws.page_margins.right=.22
    ws.page_margins.top=.35; ws.page_margins.bottom=.35
    ws.oddFooter.center.text='KJA · Planificación al 22/09/2026 · &P / &N'
    ws.oddFooter.center.size=8
def heading(ws,range_,text,size=18,color=NAVY):
    ws.merge_cells(range_); cell=ws[range_.split(':')[0]]; cell.value=text
    cell.font=Font(name='Calibri',size=size,bold=True,color=WHITE)
    cell.fill=PatternFill('solid',fgColor=color)
    cell.alignment=Alignment(vertical='center',wrap_text=True)
def cellstyle(cell,bg=None,bold=False,size=10,center=False):
    cell.font=Font(name='Calibri',size=size,color=INK,bold=bold)
    cell.alignment=Alignment(horizontal='center' if center else 'left',vertical='center',wrap_text=True)
    if bg: cell.fill=PatternFill('solid',fgColor=bg)

for ws in wb: basics(ws)

heading(read,'A1:H2','KJA | GANTT Y SPRINTS MENSUALES')
read.row_dimensions[1].height=26; read.row_dimensions[2].height=14
notes=[
('Periodo','Antecedente: 28/08/2026. Sprint 1: septiembre; Sprint 2: octubre; Sprint 3: noviembre; Sprint 4: diciembre de 2026.'),
('Corte histórico','22/09/2026. El despliegue inicial del 28/08 consta en docs/dashboard-base.md y en el commit d189536.'),
('Cómo leer el Gantt','R = día con registro Git; D = despliegue documentado; L = documento local; P = día propuesto. Los colores agrupan H.U.; las letras distinguen evidencia y planificación.'),
('Duración','Número de días marcados por tarea. En el histórico son días con evidencia, no esfuerzo ni dedicación continua. No sumar estas duraciones como duración total del proyecto.'),
('Responsables','Histórico: nombres tal como aparecen en Git, sin atribuirles responsabilidad contractual. Propuestas: roles pendientes de asignación nominal.'),
('Historias de usuario','HU-G01–HU-G12 son agrupaciones propuestas para este cronograma; no sustituyen H.U. aprobadas ni prueban que existieran en esas fechas.'),
('Todas las áreas','Se incluyen las 10 áreas de las semillas SQL. Dirección y Sistemas participan transversalmente. Confirmar catálogo vigente: no se consultó la base remota.'),
('Calendario','Propuestas de lunes a viernes; exclusiones iniciales por confirmar: 08/10, 08/12, 09/12 y 25/12. Editar Laborable en CALENDARIO. El histórico conserva registros de fines de semana.'),
('Actualizar','Editar filas existentes en TAREAS: fechas, responsable, tarea, naturaleza y evidencias. La duración y las barras se recalculan. Editar CALENDARIO modifica días propuestos.'),
('Cambiar de sprint o agregar filas','Las pestañas contienen la selección inicial por mes. Al agregar una tarea o cambiarla de mes, actualizar también su fila y fórmulas en la pestaña de destino.'),
('Límite del plan','Las fechas futuras son propuestas; no son trabajo terminado ni autorización para desplegar o publicar. El plan usa un frente principal de Sistemas y debe ajustarse a la capacidad real.'),
('Evidencia','TAREAS conserva fuente/commit, fechas de registro, RF relacionados, dependencias y criterio de salida. Registro Git no prueba un despliegue ni una prueba funcional aprobada.')]
row=4
for label,value in notes:
    read.merge_cells(start_row=row,start_column=2,end_row=row,end_column=8)
    read.cell(row,1,label); read.cell(row,2,value)
    for cell in read[row][:8]:cellstyle(cell,GRAY if row%2==0 else WHITE,bold=cell.column==1)
    read.row_dimensions[row].height=43
    row+=1
row+=1
for c,value in enumerate(['Sprint','Mes','Objetivo','Tareas','Registros','Propuestas','Inicio','Fin'],1):
    read.cell(row,c,value); cellstyle(read.cell(row,c),BLUE,True)
summary_row=row
goals=['Puesta en marcha','Consolidar y documentar','Actualizar las diez áreas','Implementar mejoras priorizadas','Validar, capacitar y cerrar']
for s in range(5):
    r=row+s+1; matching=[t for t in tasks if t['sprint']==s]
    values=['Antecedente' if s==0 else f'Sprint {s}',['Agosto','Septiembre','Octubre','Noviembre','Diciembre'][s],goals[s],
            f'=COUNTIF(\'TAREAS\'!$B$5:$B${len(tasks)+4},A{r})',
            f'=COUNTIFS(\'TAREAS\'!$B$5:$B${len(tasks)+4},A{r},\'TAREAS\'!$I$5:$I${len(tasks)+4},"<>Propuesta")',
            f'=COUNTIFS(\'TAREAS\'!$B$5:$B${len(tasks)+4},A{r},\'TAREAS\'!$I$5:$I${len(tasks)+4},"Propuesta")',
            date(2026,8,28) if s==0 else date(2026,8+s,1),date(2026,8+s,monthrange(2026,8+s)[1])]
    for c,value in enumerate(values,1):read.cell(r,c,value);cellstyle(read.cell(r,c),WHITE)
    for c in [7,8]:read.cell(r,c).number_format='dd/mm/yyyy'
    read.row_dimensions[r].height=37
for c,width in enumerate([22,18,32,12,14,14,16,16],1):read.column_dimensions[col(c)].width=width
read.freeze_panes='A4'; read.print_area=f'A1:H{summary_row+5}'

heading(source,'A1:N2','TAREAS | Fuente editable del cronograma')
source.merge_cells('A3:N3');source['A3']='Fechas futuras: propuesta. Fechas Git: registro, no inicio real ni esfuerzo. Las dependencias se revisan en la columna L.'
source.row_dimensions[3].height=25
headers=['ID','Sprint','H.U. propuesta','Tarea','Área','Responsable / autor','Inicio','Fin','Naturaleza','Días de registro (ISO)','RF / RNF','Dependencias','Criterio de salida / nota','Fuente / evidencia']
for c,label in enumerate(headers,1):source.cell(4,c,label);cellstyle(source.cell(4,c),BLUE,True)
source.row_dimensions[4].height=30
for r,t in enumerate(tasks,5):
    t['source_row']=r
    values=[t['id'],'Antecedente' if t['sprint']==0 else f"Sprint {t['sprint']}",t['hu']+' · '+H[t['hu']],t['title'],t['area'],t['owner'],
            date.fromisoformat(t['start']),date.fromisoformat(t['end']),t['kind'],', '.join(t['days']),t['rf'],t['deps'],t['criterion'],t['sources']]
    for c,value in enumerate(values,1):
        source.cell(r,c,value);cellstyle(source.cell(r,c),WHITE if r%2 else 'F5F8FB')
    source.cell(r,7).number_format=source.cell(r,8).number_format='dd/mm/yyyy'
    source.row_dimensions[r].height=78
    source.cell(r,10).comment=Comment('Fechas con evidencia separadas por coma, en formato AAAA-MM-DD. No rellenar huecos entre commits.','KJA')
widths=[10,15,49,56,25,30,15,15,25,33,29,16,66,85]
for c,width in enumerate(widths,1):source.column_dimensions[col(c)].width=width
source.freeze_panes='G5';source.auto_filter.ref=f'A4:N{len(tasks)+4}'
validation=DataValidation(type='list',formula1='"Registro Git,Despliegue documentado,Documento local,Propuesta"');validation.errorTitle='Naturaleza no válida';validation.error='Selecciona un valor de la lista.';validation.showErrorMessage=True
source.add_data_validation(validation);validation.add(f'I5:I{len(tasks)+4}')
source.print_area=f'A1:F{len(tasks)+4}';source.print_title_rows='1:4'

heading(calendar,'A1:D2','CALENDARIO | Días propuestos de trabajo')
calendar.append([])
for c,v in enumerate(['Fecha','Día','Laborable (1/0)','Nota'],1):calendar.cell(4,c,v);cellstyle(calendar.cell(4,c),BLUE,True)
days_es=['Lunes','Martes','Miércoles','Jueves','Viernes','Sábado','Domingo']
cal_rows={}
for r,d in enumerate(dates,5):
    cal_rows[d.isoformat()]=r
    values=[d,days_es[d.weekday()],cal[d.isoformat()], 'Exclusión propuesta; confirmar' if d.isoformat() in excluded else 'Fin de semana' if d.weekday()>=5 else 'Día propuesto de trabajo']
    for c,value in enumerate(values,1):calendar.cell(r,c,value);cellstyle(calendar.cell(r,c),GRAY if not cal[d.isoformat()] else WHITE)
    calendar.cell(r,1).number_format='dd/mm/yyyy'; calendar.row_dimensions[r].height=22
for c,w in enumerate([18,19,21,49],1):calendar.column_dimensions[col(c)].width=w
dv=DataValidation(type='whole',operator='between',formula1=0,formula2=1);dv.showErrorMessage=True
calendar.add_data_validation(dv);dv.add(f'C5:C{len(dates)+4}')
calendar.freeze_panes='C5';calendar.auto_filter.ref=f'A4:D{len(dates)+4}';calendar.print_area=f'A1:D{len(dates)+4}';calendar.print_title_rows='1:4'

gantt_rows={}
for sprint,ws in wsheets.items():
    month=8+sprint; first=28 if sprint==0 else 1; last=monthrange(2026,month)[1]
    monthdates=[date(2026,month,d) for d in range(first,last+1)]
    lastcol=4+len(monthdates); lastletter=col(lastcol)
    heading(ws,f'A1:{lastletter}2',('ANTECEDENTE · AGOSTO' if sprint==0 else f'SPRINT {sprint} · '+['','SEPTIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE'][sprint])+' 2026 | GANTT')
    ws.row_dimensions[1].height=24;ws.row_dimensions[2].height=16
    ws.merge_cells(f'A3:{lastletter}3');ws['A3']=goals[sprint]+' · Corte histórico: 22/09/2026 · Fechas futuras propuestas'
    ws.row_dimensions[3].height=23
    ws.merge_cells(f'A4:{lastletter}4');ws['A4']='R registro Git   ·   D despliegue documentado   ·   L documento local   ·   P propuesta   |   Duración = días marcados, no horas de esfuerzo'
    ws.row_dimensions[4].height=23
    ws.merge_cells(f'A5:{lastletter}5');ws['A5']='Editar TAREAS para actualizar fechas y contenido. HU-G son agrupaciones propuestas; responsables futuros pendientes de confirmación.'
    ws.row_dimensions[5].height=22
    labels=['Historia de usuario propuesta','Tarea / entregable','Responsable / autor','Días']
    for c,label in enumerate(labels,1):
        ws.merge_cells(start_row=7,start_column=c,end_row=8,end_column=c)
        ws.cell(7,c,label);cellstyle(ws.cell(7,c),BLUE,True)
    for c,d in enumerate(monthdates,5):
        ws.cell(7,c,d);ws.cell(7,c).number_format='dd/mm'
        ws.cell(8,c,['L','M','X','J','V','S','D'][d.weekday()])
        for r in [7,8]:cellstyle(ws.cell(r,c),BLUE if cal[d.isoformat()] else 'DEE1E5',True,9,True)
        ws.column_dimensions[col(c)].width=6
    ws.row_dimensions[7].height=26;ws.row_dimensions[8].height=19
    matching=[t for t in tasks if t['sprint']==sprint]
    gantt_rows[sprint]=[]
    for r,t in enumerate(matching,9):
        sr=t['source_row']; color=COLORS[(int(t['hu'][-2:])-1)%len(COLORS)]
        ws.cell(r,1,f"='TAREAS'!C{sr}");ws.cell(r,2,f"='TAREAS'!A{sr}&\" · \"&'TAREAS'!D{sr}")
        ws.cell(r,3,f"='TAREAS'!F{sr}")
        ws.cell(r,4,f'=COUNTIF(E{r}:{lastletter}{r},"R")+COUNTIF(E{r}:{lastletter}{r},"D")+COUNTIF(E{r}:{lastletter}{r},"L")+COUNTIF(E{r}:{lastletter}{r},"P")')
        for c in range(1,5):cellstyle(ws.cell(r,c),color if c==1 else WHITE,c==1,10,c==4)
        ws.cell(r,4).number_format='0" días"'
        ws.cell(r,2).comment=Comment(t['criterion']+'\n\nFuente: '+t['sources']+'\nDependencias: '+(t['deps'] or 'Sin dependencia registrada'),'KJA')
        for c,d in enumerate(monthdates,5):
            letter=col(c);cr=cal_rows[d.isoformat()]
            ws.cell(r,c,f'=IF(\'TAREAS\'!$I${sr}="Propuesta",IF(AND({letter}$7>=\'TAREAS\'!$G${sr},{letter}$7<=\'TAREAS\'!$H${sr},\'CALENDARIO\'!$C${cr}=1),"P",""),IF(ISNUMBER(SEARCH(TEXT({letter}$7,"yyyy-mm-dd"),\'TAREAS\'!$J${sr})),IF(\'TAREAS\'!$I${sr}="Registro Git","R",IF(\'TAREAS\'!$I${sr}="Despliegue documentado","D","L")),""))')
            cellstyle(ws.cell(r,c),WHITE if cal[d.isoformat()] else GRAY,True,10,True)
        ws.conditional_formatting.add(f'E{r}:{lastletter}{r}',FormulaRule(formula=[f'E{r}<>""'],fill=PatternFill('solid',fgColor=color,bgColor=color)))
        for c in range(1,lastcol+1):ws.cell(r,c).border=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)
        ws.row_dimensions[r].height=72
        gantt_rows[sprint].append((r,t['id']))
    for c,width in enumerate([39,59,28,11],1):ws.column_dimensions[col(c)].width=width
    ws.freeze_panes='E9';ws.print_title_rows='1:8';ws.print_title_cols='A:D'
    ws.print_area=f'A1:{lastletter}{8+len(matching)}'
    ws.sheet_properties.tabColor=COLORS[sprint%len(COLORS)]
    ws.sheet_view.zoomScale=70

# Pestañas de presentación primero y fuentes al final.
wb._sheets=[read,*wsheets.values(),source,calendar]
wb.active=2
file=OUT/'KJA_Gantt_Sprints_Septiembre_Diciembre_2026.xlsx'
wb.save(file)
metadata={'cutoff':CUT.isoformat(),'tasks':tasks,'histories':H,'excluded_dates_proposed':sorted(excluded),
          'gantt_rows':gantt_rows,'summary_row':summary_row,'output':str(file)}
(OUT/'gantt-datos.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')

md='''# Gantt y sprints mensuales — Portal KJA

**Corte:** 22/09/2026 · **Horizonte:** hasta diciembre de 2026 · **Estado:** histórico documentado y propuesta futura.

El primer despliegue consta el **28/08/2026** en `docs/dashboard-base.md` y en el
commit `d189536`. Se muestra como antecedente y no se desplaza a septiembre.
Sprint 1 corresponde a septiembre; Sprint 2 a octubre; Sprint 3 a noviembre;
Sprint 4 a diciembre. Todos abarcan el mes calendario completo.

[Abrir el Gantt editable en Excel](../outputs/gantt-kja-20260922/KJA_Gantt_Sprints_Septiembre_Diciembre_2026.xlsx).

## Lectura y supuestos

- R: registro Git; D: despliegue declarado en una guía; L: documento local; P: propuesta futura.
- Se marcan únicamente fechas de evidencia histórica; los días entre commits no se interpretan como trabajo continuo. Las fechas Git usan el offset -05:00 registrado.
- Duración significa días marcados por tarea; no es esfuerzo, tiempo exclusivo ni duración total del proyecto.
- Los autores Git se conservan literalmente. Los responsables de propuestas son roles pendientes de asignación; no se atribuyen tareas a Dany o Yeiser por aparecer en la imagen de referencia.
- HU-G01–HU-G12 son agrupaciones propuestas para el cronograma, no historias aprobadas. Se mantienen referencias RF/RNF y se señala cuando falta formalizar un requisito.
- Propuestas de lunes a viernes. Se excluyen inicialmente 08/10, 08/12, 09/12 y 25/12 como fechas no laborables **por confirmar**; no se afirma un calendario laboral oficial. La hoja CALENDARIO permite modificarlas.
- El plan propone un frente principal secuencial de Sistemas. Fechas, disponibilidad de líderes, alcance y capacidad requieren confirmación antes de comprometer entregas.
- Las diez áreas se obtienen de `supabase/asistencia_schema.sql` y `supabase/asistencia_seed_colaboradores.sql`; no se certifica el catálogo remoto actual.
- Clínica y Reclutamiento se limitan a la operación laboral del portal. No se amplía el alcance a expedientes clínicos, datos de pacientes o expedientes de candidatos; Contabilidad no incorpora planillas o pagos.
- La validación de Marketing no autoriza publicaciones reales. Las tareas futuras no se presentan como ejecutadas.

## Objetivos y revisión mensual

| Sprint | Periodo | Objetivo | Revisión propuesta |
|---|---|---|---|
| Antecedente | 28–31/08/2026 | Conservar evidencia del despliegue inicial | No aplica retrospectiva inventada |
| Sprint 1 | 01–30/09/2026 | Consolidar portal, documentar y acordar prioridades | 30/09/2026 |
| Sprint 2 | 01–31/10/2026 | Revisar y actualizar las diez áreas | 30/10/2026 |
| Sprint 3 | 01–30/11/2026 | Implementar y verificar mejoras priorizadas | 30/11/2026 |
| Sprint 4 | 01–31/12/2026 | Validar por área, capacitar y cerrar el año | 31/12/2026 |

## Tareas y trazabilidad

'''
for sprint in range(5):
    md+=f'### {"Antecedente" if sprint==0 else "Sprint "+str(sprint)}\n\n| ID | H.U. propuesta | Tarea | Inicio–fin | Días marcados | Naturaleza | Responsable / autor | Dependencia |\n|---|---|---|---|---:|---|---|---|\n'
    for t in tasks:
        if t['sprint']==sprint:md+=f"| {t['id']} | {t['hu']} | {t['title']} | {t['start']} → {t['end']} | {t['duration']} | {t['kind']} | {t['owner']} | {t['deps'] or '—'} |\n"
    md+='\n'
md+='## Historias propuestas para agrupar el Gantt\n\n| ID provisional | Historia |\n|---|---|\n'
for key,value in H.items():md+=f'| {key} | {value} |\n'
md+='\n## Evidencias y criterios de salida\n\n'
for t in tasks:
    md+=f"- **{t['id']}** · {t['rf'] or 'Por vincular'}. {t['criterion']} Fuente: {t['sources']}.\n"
md+='''
## Actualización

El Excel contiene una pestaña por sprint, la tabla editable TAREAS y el CALENDARIO.
Las fórmulas actualizan fechas, texto, responsables y marcas de las filas existentes.
Agregar tareas o cambiarlas de mes requiere actualizar su fila en el sprint de
destino. El plan documental de este archivo es la línea base; tras editar Excel,
registrar los cambios aprobados en la documentación para evitar versiones divergentes.
'''
(ROOT/'docs'/'28-gantt-sprints-mensuales.md').write_text(md,encoding='utf-8',newline='\n')
print(json.dumps({'file':str(file),'tasks':len(tasks),'sheets':wb.sheetnames,'by_sprint':dict(Counter(t['sprint'] for t in tasks))},ensure_ascii=False))
