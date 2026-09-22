from pathlib import Path
import json, time
import win32com.client
from openpyxl import load_workbook
import pymupdf as fitz

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'outputs'/'gantt-kja-20260922'
meta=json.loads((OUT/'gantt-datos.json').read_text(encoding='utf-8'))
file=Path(meta['output'])
qa=OUT/'revision';qa.mkdir(exist_ok=True)
app=None;book=None
try:
    app=win32com.client.DispatchEx('Excel.Application')
    app.Visible=False;app.DisplayAlerts=False;app.EnableEvents=False
    app.AutomationSecurity=3
    print('Excel',app.Version,'archivo',file.exists(),'bytes',file.stat().st_size,flush=True)
    print('Openpyxl:',len(load_workbook(file).sheetnames),'hojas válidas',flush=True)
    book=app.Workbooks.Open(str(file),0,False)
    app.CalculateFullRebuild()
    for _ in range(40):
        if app.CalculationState==0:break
        time.sleep(.25)
    assert app.CalculationState==0,'Excel no terminó el cálculo'
    # Comprobar reactividad, dejando intactos los datos finales.
    sample=next(t for t in meta['tasks'] if t['sprint']==2)
    sr=sample['source_row']
    cell=book.Worksheets('TAREAS').Cells(sr,8)
    original=cell.Value2
    cell.Value2=book.Worksheets('TAREAS').Cells(sr,7).Value2
    app.CalculateFullRebuild()
    assert book.Worksheets('SPRINT 2').Cells(9,4).Value2==1,'No actualizó duración al cambiar fecha'
    cell.Value2=original
    app.CalculateFullRebuild()
    assert book.Worksheets('SPRINT 2').Cells(9,4).Value2==2,'No restauró duración'
    # Reacción al calendario editable, sin alterar el histórico.
    cal=book.Worksheets('CALENDARIO')
    target=next(r for r in range(5,cal.UsedRange.Rows.Count+1) if cal.Cells(r,1).Text=='01/10/2026')
    old=cal.Cells(target,3).Value2;cal.Cells(target,3).Value2=0
    app.CalculateFullRebuild()
    assert book.Worksheets('SPRINT 2').Cells(9,4).Value2==1,'Calendario no actualiza Gantt'
    cal.Cells(target,3).Value2=old
    app.CalculateFullRebuild()
    colored=book.Worksheets('SPRINT 2').Cells(9,5)
    print('Color visible de barra:',colored.DisplayFormat.Interior.Color,flush=True)
    assert colored.DisplayFormat.Interior.Color!=16777215,'La barra del Gantt no muestra su color'
    book.Save()
    for sheet in book.Worksheets:
        sheet.ExportAsFixedFormat(0,str(qa/(sheet.Name.replace(' ','_')+'.pdf')))
    book.Close(SaveChanges=False);book=None
finally:
    if book is not None:book.Close(SaveChanges=False)
    if app is not None:app.Quit()

values=load_workbook(file,data_only=True)
formulas=load_workbook(file,data_only=False)
errors=[];formula_count=0
for ws in formulas:
    for row in ws:
        for cell in row:
            cached=values[ws.title][cell.coordinate]
            if cached.data_type=='e':errors.append((ws.title,cell.coordinate,cached.value))
            if cell.data_type=='f':
                formula_count+=1
                if cached.value is None and cell.column<=4:errors.append((ws.title,cell.coordinate,'Sin resultado de fórmula'))
assert not errors,errors[:10]
taskmap={t['id']:t for t in meta['tasks']}
for sprint,rows in meta['gantt_rows'].items():
    name='ANTECEDENTE' if sprint=='0' else 'SPRINT '+sprint
    sheet=values[name]
    for row,ident in rows:
        t=taskmap[ident]
        actual=[sheet.cell(7,c).value.date().isoformat() for c in range(5,sheet.max_column+1) if sheet.cell(row,c).value]
        assert actual==t['marked_days'],(ident,actual,t['marked_days'])
        assert sheet.cell(row,4).value==t['duration'],(ident,'duración incorrecta')
        expected={'Registro Git':'R','Despliegue documentado':'D','Documento local':'L','Propuesta':'P'}[t['kind']]
        assert all(sheet.cell(row,c).value in (None,'',expected) for c in range(5,sheet.max_column+1)),ident
for sprint in range(5):
    row=meta['summary_row']+sprint+1
    assert values['LEEME'].cell(row,4).value==sum(t['sprint']==sprint for t in meta['tasks'])

renders=[]
for pdf in sorted(qa.glob('*.pdf')):
    document=fitz.open(pdf)
    assert len(document)>0,pdf
    for n,page in enumerate(document):
        assert len(page.get_text().strip())>30,(pdf,n,'Página vacía')
    first=document[0]
    target=qa/(pdf.stem+'.png')
    first.get_pixmap(matrix=fitz.Matrix(1.25,1.25)).save(target)
    renders.append({'sheet':pdf.stem,'pages':len(document),'preview':str(target)})
    if pdf.stem=='SPRINT_1' and len(document)>1:
        document[-1].get_pixmap(matrix=fitz.Matrix(1.25,1.25)).save(qa/'SPRINT_1_final.png')
report={'tasks_checked':len(meta['tasks']),'formulas_checked':formula_count,'formula_errors':errors,
        'date_and_calendar_edit_checks':'passed','renders':renders}
(qa/'verificacion.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
